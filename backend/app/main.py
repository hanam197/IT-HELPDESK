import csv
import io
import json
import secrets
from pathlib import Path
from datetime import date, datetime, timedelta, timezone
from collections import Counter
import ipaddress
import jwt
import qrcode
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File, Query
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse
from fastapi.encoders import jsonable_encoder
from pydantic import ValidationError
from sqlalchemy import select, func, or_, cast, String, inspect
from sqlalchemy.exc import IntegrityError
from openpyxl import Workbook, load_workbook
from PIL import Image, UnidentifiedImageError
from .database import get_db, settings
from . import models as m
from .schemas import RESOURCES, READ_ONLY, SCHEMAS, Login, Move, Assign, Return, Transfer, Comment
from .security import current_user, authorize, passwords, check_login_limit
from .services import get, serialize, save, audit, activity, move, assign, return_asset, master, next_number

app = FastAPI(title='IT Helpdesk & Asset Management', version='1.0.0')

@app.middleware('http')
async def security_headers(request, call_next):
    if request.method not in {'GET','HEAD','OPTIONS'} and request.url.path.startswith('/api/'):
        # Custom header prevents cross-origin HTML form CSRF; no cross-origin CORS is allowed.
        if request.headers.get('x-requested-with') != 'Helpdesk':
            return JSONResponse({'detail':'Missing request verification header'},403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Cache-Control'] = 'no-store'
    return response

@app.exception_handler(IntegrityError)
async def integrity_handler(request, exc):
    return JSONResponse({'detail':'Duplicate value or invalid relationship. Check asset code, serial, IP, MAC and active records.'},409)

@app.get('/api/health')
def health(db=Depends(get_db)):
    db.execute(select(1)); return {'status':'ok'}

@app.post('/api/auth/login')
def login(payload: Login, request: Request, response: Response, db=Depends(get_db)):
    check_login_limit((request.client.host if request.client else 'local',payload.username))
    user = db.scalar(select(m.User).where(m.User.username == payload.username,m.User.archived == False))
    dummy = '$argon2id$v=19$m=65536,t=3,p=4$UVNhbHRGb3JEZW1vMTIzNA$AAECAwQFBgcICQoLDA0ODxAREhMUFRYXGBkaGxwdHh8'
    try: valid = passwords.verify(payload.password, user.password_hash if user else dummy)
    except Exception: valid = False
    if not user or not valid: raise HTTPException(401,'Incorrect username or password')
    token = jwt.encode({'sub':str(user.id),'exp':datetime.now(timezone.utc)+timedelta(hours=8)},settings.secret_key,algorithm='HS256')
    response.set_cookie('helpdesk_session',token,httponly=True,secure=settings.secure_cookie,samesite='strict',max_age=28800)
    return serialize(user)

@app.post('/api/auth/logout')
def logout(response: Response):
    response.delete_cookie('helpdesk_session'); return {'ok':True}

@app.get('/api/auth/me')
def me(user=Depends(current_user)): return serialize(user)

@app.get('/api/meta')
def metadata(user=Depends(current_user),db=Depends(get_db)):
    result={}
    for key in ['users','master-data','asset-types','locations','assets','warehouses','inventory-items','vlans','subnets','interfaces','articles']:
        rows=db.scalars(select(RESOURCES[key]).where(RESOURCES[key].archived == False)).all()
        result[key]=[({'id':r.id,'name':r.name,'username':r.username,'department':r.department} if key=='users' else ({**serialize(r),'path':location_path(db,r.id)} if key=='locations' else serialize(r))) for r in rows]
    return result

def location_path(db,id):
    names=[]; seen=set()
    while id and id not in seen:
        seen.add(id); loc=db.get(m.Location,id)
        if not loc: break
        names.insert(0,loc.name); id=loc.parent_id
    return ' / '.join(names)

def enriched(db,obj):
    data=serialize(obj)
    for c in inspect(type(obj)).columns:
        if c.name == 'password_hash': continue
        for fk in c.foreign_keys:
            value=getattr(obj,c.name)
            target=next((x for x in RESOURCES.values() if x.__tablename__==fk.column.table.name),None)
            linked=db.get(target,value) if value and target else None
            if linked:
                label=getattr(linked,'code',None) or getattr(linked,'name',None) or getattr(linked,'number',None) or getattr(linked,'cidr',None) or str(linked.id)
                if isinstance(linked,m.MasterData): label=linked.name
                data[c.name.removesuffix('_id')+'_label']=label
    if isinstance(obj,m.Asset):
        loc=db.scalar(select(m.LocationHistory).where(m.LocationHistory.asset_id==obj.id,m.LocationHistory.ended_at==None))
        ass=db.scalar(select(m.Assignment).where(m.Assignment.asset_id==obj.id,m.Assignment.returned_at==None))
        data['location_label']=location_path(db,loc.location_id) if loc else 'Unlocated'
        data['location_id']=loc.location_id if loc else None
        data['assigned_to']=db.get(m.User,ass.user_id).name if ass else None
        data['assignment_id']=ass.id if ass else None
        data['assignment_user_id']=ass.user_id if ass else None
        data['current_assignment_since']=ass.created_at.isoformat() if ass else None
        data['location_since']=loc.created_at.isoformat() if loc else None
        data['since']=ass.created_at.isoformat() if ass else (loc.created_at.isoformat() if loc else None)
        data['photo_url']=f'/api/assets/{obj.id}/photo' if obj.photo else None
    if isinstance(obj,m.AssetOperation):
        for prefix in ('from', 'to'):
            entity_type=getattr(obj, f'{prefix}_entity_type')
            entity_id=getattr(obj, f'{prefix}_entity_id')
            label_value=None
            if entity_type == 'USER' and entity_id:
                linked=db.get(m.User, entity_id); label_value=linked.name if linked else None
            elif entity_type == 'LOCATION' and entity_id:
                label_value=location_path(db, entity_id)
            elif entity_type == 'SUPPLIER':
                label_value='Supplier'
            data[f'{prefix}_label']=label_value
    if isinstance(obj,m.IPAddress):
        interface=db.get(m.NetworkInterface,obj.interface_id) if obj.interface_id else None
        subnet=db.get(m.Subnet,obj.subnet_id)
        data['vlan_label']=db.get(m.VLAN,subnet.vlan_id).name
        if interface:
            asset=db.get(m.Asset,interface.asset_id); a=enriched(db,asset)
            data.update(asset_id=asset.id,asset_label=asset.code,mac=interface.mac,hostname=interface.hostname,location_label=a['location_label'])
            port=db.scalar(select(m.SwitchPort).where(m.SwitchPort.connected_asset_id==asset.id))
            if port: data.update(switch_label=db.get(m.Asset,port.switch_id).code,switch_name=db.get(m.Asset,port.switch_id).name,port=port.name)
    if isinstance(obj,m.SwitchPort):
        data['tagged_vlans_label']=', '.join(str(db.get(m.VLAN,id).tag) for id in obj.tagged_vlans if db.get(m.VLAN,id))
    if isinstance(obj,m.Subnet):
        network=ipaddress.ip_network(obj.cidr); total=network.num_addresses-(2 if network.version==4 and network.prefixlen<31 else 0)
        used=db.scalar(select(func.count()).select_from(m.IPAddress).where(m.IPAddress.subnet_id==obj.id,m.IPAddress.interface_id!=None,m.IPAddress.archived==False))
        data.update(total_ips=total,used_ips=used,available_ips=total-used)
    return data

def query_rows(db,resource,user,q='',filters='{}',view='',location=''):
    if resource not in RESOURCES: raise HTTPException(404,'Unknown resource')
    if resource=='users' and user.role!='ADMIN': raise HTTPException(403,'Admin only')
    model=RESOURCES[resource]; stmt=select(model)
    if hasattr(model,'archived'): stmt=stmt.where(model.archived==False)
    if resource=='ticket-activities' and user.role=='VIEWER': stmt=stmt.where(m.TicketActivity.internal==False)
    try: values=json.loads(filters)
    except ValueError: raise HTTPException(422,'Invalid filters')
    if not isinstance(values,dict): raise HTTPException(422,'Filters must be an object')
    for key,value in values.items():
        if key not in inspect(model).columns or key=='password_hash': raise HTTPException(422,f'Invalid filter: {key}')
        col=getattr(model,key); stmt=stmt.where(col==value)
    if resource=='tickets':
        if view in {'open','overdue'}:
            closed=select(m.MasterData.id).where(m.MasterData.group=='ticket_status',m.MasterData.code.in_(['resolved','closed','cancelled']))
            stmt=stmt.where(m.Ticket.status_id.not_in(closed))
        if view=='overdue': stmt=stmt.where(m.Ticket.due_at<m.now())
        if view=='resolved-today': stmt=stmt.where(m.Ticket.resolved_at>=datetime.now(timezone.utc).replace(hour=0,minute=0,second=0,microsecond=0))
    if resource=='assignments' and view=='returned': stmt=stmt.where(m.Assignment.returned_at!=None)
    if resource=='assets' and location:
        ids=[r.id for r in db.scalars(select(m.Location)) if location_path(db,r.id)==location]
        stmt=stmt.where(m.Asset.id.in_(select(m.LocationHistory.asset_id).where(m.LocationHistory.location_id.in_(ids),m.LocationHistory.ended_at==None)))
    if q:
        cols=[c for c in inspect(model).columns if (isinstance(c.type,String) or c.name=='created_at') and c.name!='password_hash']
        conditions=[cast(c,String).ilike('%'+q+'%') for c in cols]
        if resource=='ip-addresses':
            pattern='%'+q+'%'
            matching_assets=select(m.Asset.id).where(or_(m.Asset.code.ilike(pattern),m.Asset.name.ilike(pattern),m.Asset.serial.ilike(pattern)))
            matching_switches=select(m.Asset.id).where(or_(m.Asset.code.ilike(pattern),m.Asset.name.ilike(pattern)))
            connected=select(m.SwitchPort.connected_asset_id).where(or_(m.SwitchPort.name.ilike(pattern),m.SwitchPort.switch_id.in_(matching_switches)))
            interfaces=select(m.NetworkInterface.id).where(or_(m.NetworkInterface.mac.ilike(pattern),m.NetworkInterface.hostname.ilike(pattern),m.NetworkInterface.asset_id.in_(matching_assets),m.NetworkInterface.asset_id.in_(connected)))
            conditions.append(m.IPAddress.interface_id.in_(interfaces))
        stmt=stmt.where(or_(*conditions))
    return stmt

@app.get('/api/dashboard')
def dashboard(user=Depends(current_user),db=Depends(get_db)):
    assets=[enriched(db,r) for r in db.scalars(select(m.Asset).where(m.Asset.archived==False))]
    maintenance=[enriched(db,r) for r in db.scalars(select(m.Maintenance).where(m.Maintenance.archived==False))]
    codes={r.id:r.code for r in db.scalars(select(m.MasterData))}
    active=[x for x in maintenance if x['end_at'] is None]
    stats={'Total assets':len(assets),'Available assets':sum(codes[a['status_id']]=='available' for a in assets),'Active assets':sum(codes[a['status_id']] in {'active','in_use'} for a in assets),'Assigned assets':sum(bool(a['assignment_id']) for a in assets),'Assets in warehouse':sum(codes[a['status_id']]=='available' and 'warehouse' in (a.get('location_label') or '').lower() for a in assets),'Assets under repair':sum(codes[a['status_id']] in {'repair','maintenance'} for a in assets),'Broken assets':sum(codes[a['status_id']]=='broken' for a in assets),'Low stock items':sum(float(r.quantity)<float(r.minimum_stock) for r in db.scalars(select(m.InventoryItem).where(m.InventoryItem.archived==False))),'Active maintenance':len(active)}
    distribution=lambda rows,key:dict(Counter(r.get(key) or 'Unassigned' for r in rows))
    attention=[{'title':a['code']+' · '+a['name'],'subtitle':'Asset under repair','resource':'assets','id':a['id']} for a in assets if codes[a['status_id']] in {'repair','maintenance','broken'}]
    attention += [{'title':x['number'],'subtitle':'Active maintenance','resource':'maintenance','id':x['id']} for x in active]
    for ip in db.scalars(select(m.IPAddress).where(m.IPAddress.archived==False)):
        if codes[ip.status_id]=='conflict': attention.append({'title':ip.address,'subtitle':'IP conflict','resource':'ip-addresses','id':ip.id})
    operations=[enriched(db,r) for r in db.scalars(select(m.AssetOperation).order_by(m.AssetOperation.operation_date.desc()).limit(12))]
    return {'stats':stats,'asset_type':distribution(assets,'type_label'),'asset_status':distribution(assets,'status_label'),'asset_location':distribution(assets,'location_label'),'operations':operations,'recent_returns':[r for r in operations if r['operation_type']=='RETURN'][:5],'recent_assignments':[r for r in operations if r['operation_type']=='ASSIGN'][:5],'attention':attention,'activities':[enriched(db,a) for a in db.scalars(select(m.AuditLog).order_by(m.AuditLog.id.desc()).limit(8))]}

@app.get('/api/search')
def search(q:str=Query(min_length=2,max_length=150),user=Depends(current_user),db=Depends(get_db)):
    result=[]; asset_ids=set()
    for resource in ['assets','interfaces','ip-addresses','locations','articles','users','warehouses','inventory-items']:
        model=RESOURCES[resource]
        cols=[c for c in inspect(model).columns if isinstance(c.type,String) and c.name!='password_hash']
        rows=db.scalars(select(model).where(model.archived==False,or_(*[cast(c,String).ilike('%'+q+'%') for c in cols])).limit(15))
        for row in rows:
            data=enriched(db,row)
            if resource=='users': data={k:data[k] for k in ['id','name','username','department']}
            result.append({'resource':resource,**data})
            if resource=='interfaces': asset_ids.add(row.asset_id)
            if resource=='ip-addresses' and row.interface_id: asset_ids.add(db.get(m.NetworkInterface,row.interface_id).asset_id)
            if resource=='assets': asset_ids.add(row.id)
    found={r['id'] for r in result if r['resource']=='assets'}
    for id in asset_ids-found: result.append({'resource':'assets',**enriched(db,db.get(m.Asset,id))})
    return result

@app.get('/api/assets/{id}/detail')
def asset_detail(id:int,user=Depends(current_user),db=Depends(get_db)):
    asset=get(db,m.Asset,id); data=enriched(db,asset)
    data['asset_type']=serialize(get(db,m.AssetType,asset.type_id))
    for resource,model,field in [('location-history',m.LocationHistory,'asset_id'),('assignments',m.Assignment,'asset_id'),('asset-operations',m.AssetOperation,'asset_id'),('interfaces',m.NetworkInterface,'asset_id'),('maintenance',m.Maintenance,'asset_id'),('switch-ports',m.SwitchPort,'switch_id')]:
        data[resource]=[enriched(db,r) for r in db.scalars(select(model).where(getattr(model,field)==id).order_by(model.id.desc()))]
    ids=[r['id'] for r in data['interfaces']]
    data['ip-addresses']=[enriched(db,r) for r in db.scalars(select(m.IPAddress).where(m.IPAddress.interface_id.in_(ids)))]
    data['connected_ports']=[enriched(db,r) for r in db.scalars(select(m.SwitchPort).where(m.SwitchPort.connected_asset_id==id))]
    data['audit-logs']=[enriched(db,r) for r in db.scalars(select(m.AuditLog).where(or_((m.AuditLog.object_type=='assets')&(m.AuditLog.object_id==id),m.AuditLog.new_value['asset_id'].as_integer()==id)).order_by(m.AuditLog.id.desc()))]
    return data

@app.get('/api/locations/{id}/detail')
def location_detail(id:int,user=Depends(current_user),db=Depends(get_db)):
    location=get(db,m.Location,id); data=enriched(db,location); data['path']=location_path(db,id)
    current_asset_ids=select(m.LocationHistory.asset_id).where(m.LocationHistory.location_id==id,m.LocationHistory.ended_at==None)
    data['assets']=[enriched(db,a) for a in db.scalars(select(m.Asset).where(m.Asset.id.in_(current_asset_ids),m.Asset.archived==False).order_by(m.Asset.code))]
    return data

@app.post('/api/assets/{id}/move')
def move_asset(id:int,payload:Move,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'operations'); row=move(db,id,payload,user); db.commit(); return serialize(row)
@app.post('/api/assets/{id}/assign')
def assign_asset(id:int,payload:Assign,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'operations'); row=assign(db,id,payload,user); db.commit(); return serialize(row)
@app.post('/api/assets/{id}/return')
def return_device(id:int,payload:Return,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'operations'); row=return_asset(db,id,payload,user); db.commit(); return serialize(row)
@app.post('/api/assets/{id}/transfer')
def transfer_device(id:int,payload:Transfer,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'operations'); return_asset(db,id,Return(condition_in=payload.condition_in,note=payload.note),user)
    row=assign(db,id,Assign(**payload.model_dump(exclude={'condition_in'})),user); db.commit(); return serialize(row)

@app.get('/api/assets/{id}/qr')
def asset_qr(id:int,user=Depends(current_user),db=Depends(get_db)):
    get(db,m.Asset,id); buffer=io.BytesIO(); qrcode.make(f'{settings.frontend_url}/assets/{id}').save(buffer,format='PNG'); buffer.seek(0)
    return StreamingResponse(buffer,media_type='image/png')

@app.post('/api/assets/register',status_code=201)
def register_asset(payload:dict,user=Depends(current_user),db=Depends(get_db)):
    """Create an asset and its first immutable location record atomically."""
    authorize(user,'assets')
    location_id=payload.pop('location_id',None)
    if not location_id: raise HTTPException(422,'Location is required when creating an asset')
    asset=save(db,'assets',parse_payload('assets',payload),user)
    move(db,asset.id,Move(location_id=location_id,reason='Initial asset registration'),user)
    db.commit(); return enriched(db,asset)

@app.post('/api/inventory/transactions',status_code=201)
def inventory_transaction(payload:dict,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'warehouses')
    transaction_type=str(payload.get('transaction_type') or '').upper()
    if transaction_type not in {'RECEIVE','ISSUE'}: raise HTTPException(422,'Transaction type must be RECEIVE or ISSUE')
    warehouse_id=payload.get('warehouse_id'); item_id=payload.get('item_id'); quantity=payload.get('quantity')
    if not warehouse_id or not item_id or quantity is None: raise HTTPException(422,'Warehouse, item and quantity are required')
    try: quantity=float(quantity)
    except (TypeError,ValueError): raise HTTPException(422,'Quantity must be numeric')
    if quantity <= 0: raise HTTPException(422,'Quantity must be greater than zero')
    warehouse=get(db,m.Warehouse,warehouse_id)
    item=db.scalar(select(m.InventoryItem).where(m.InventoryItem.id==item_id,m.InventoryItem.archived==False).with_for_update())
    if not item: raise HTTPException(404,'Inventory item not found')
    if item.warehouse_id != warehouse.id: raise HTTPException(422,'Item does not belong to this warehouse')
    current=float(item.quantity or 0)
    if transaction_type=='ISSUE' and current < quantity: raise HTTPException(422,'Insufficient stock')
    item.quantity=current + quantity if transaction_type=='RECEIVE' else current - quantity
    row=m.InventoryTransaction(number=next_number(db,'STK'),transaction_type=transaction_type,warehouse_id=warehouse.id,item_id=item.id,quantity=quantity,source_vendor=payload.get('source_vendor'),condition=payload.get('condition'),performed_by=user.id,note=payload.get('note'))
    db.add(row); db.flush(); audit(db,user,transaction_type.lower(),'inventory-transactions',row); db.commit()
    return enriched(db,row)

@app.post('/api/assets/{id}/photo')
async def upload_asset_photo(id:int,file:UploadFile=File(...),user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'assets'); asset=get(db,m.Asset,id)
    content=await file.read(5*1024*1024+1)
    if len(content)>5*1024*1024: raise HTTPException(413,'Maximum photo size is 5 MB')
    try:
        image=Image.open(io.BytesIO(content)); image.verify(); image_format=image.format
    except (UnidentifiedImageError,OSError): raise HTTPException(422,'Select a valid JPG, PNG or WebP image')
    extensions={'JPEG':'jpg','PNG':'png','WEBP':'webp'}
    if image_format not in extensions: raise HTTPException(422,'Only JPG, PNG and WebP photos are supported')
    directory=Path(settings.upload_dir)/'asset-photos'; directory.mkdir(parents=True,exist_ok=True)
    key=f'{id}-{secrets.token_hex(16)}.{extensions[image_format]}'; path=directory/key; path.write_bytes(content)
    old=serialize(asset); previous=asset.photo; asset.photo=key
    audit(db,user,'photo uploaded','assets',asset,old); db.commit()
    if previous: (directory/Path(previous).name).unlink(missing_ok=True)
    return enriched(db,asset)

@app.get('/api/assets/{id}/photo')
def asset_photo(id:int,user=Depends(current_user),db=Depends(get_db)):
    asset=get(db,m.Asset,id)
    if not asset.photo: raise HTTPException(404,'Asset has no photo')
    path=Path(settings.upload_dir)/'asset-photos'/Path(asset.photo).name
    if not path.is_file(): raise HTTPException(404,'Photo file not found')
    return FileResponse(path)

@app.get('/api/assets/import/template')
def asset_import_template(user=Depends(current_user)):
    book=Workbook(); sheet=book.active; sheet.title='Assets'
    sheet.append(['asset_name','asset_type','status','location','brand','model','serial','received_date','handover_date','description'])
    sheet.append(['Example laptop','Laptop','Available','Q7 / OUTBOUND / DG-01BD','Dell','Latitude 5440','SN-EXAMPLE-001','','',''])
    sheet.freeze_panes='A2'; buffer=io.BytesIO(); book.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="asset-import-template.xlsx"'})

@app.post('/api/assets/import',status_code=201)
async def import_assets(file:UploadFile=File(...),user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'assets'); content=await file.read(5*1024*1024+1)
    if len(content)>5*1024*1024: raise HTTPException(413,'Maximum import size is 5 MB')
    try: book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    except Exception: raise HTTPException(422,'Select a valid XLSX workbook')
    rows=book.active.iter_rows(values_only=True); headers=[str(v or '').strip().lower() for v in next(rows,())]
    required={'asset_name','asset_type','status','location','model','serial'}
    if not required.issubset(headers): raise HTTPException(422,f"Missing columns: {', '.join(sorted(required-set(headers)))}")
    created=[]
    try:
        for row_number,values in enumerate(rows,start=2):
            record=dict(zip(headers,values))
            if not any(v not in (None,'') for v in values): continue
            def required_text(key):
                value=str(record.get(key) or '').strip()
                if not value: raise HTTPException(422,f'Row {row_number}: {key} is required')
                return value
            type_name=required_text('asset_type'); status_value=required_text('status'); location_value=required_text('location')
            asset_type=db.scalar(select(m.AssetType).where(func.lower(m.AssetType.name)==type_name.lower(),m.AssetType.archived==False))
            status=db.scalar(select(m.MasterData).where(m.MasterData.group=='asset_status',m.MasterData.archived==False,or_(func.lower(m.MasterData.name)==status_value.lower(),func.lower(m.MasterData.code)==status_value.lower().replace(' ','_'))))
            locations=[loc for loc in db.scalars(select(m.Location).where(m.Location.archived==False)) if location_path(db,loc.id).lower()==location_value.lower()]
            if not asset_type: raise HTTPException(422,f'Row {row_number}: unknown asset type "{type_name}"')
            if not status: raise HTTPException(422,f'Row {row_number}: unknown asset status "{status_value}"')
            if len(locations)!=1: raise HTTPException(422,f'Row {row_number}: location must match one full path, for example Q7 / OFFICE')
            data={'name':required_text('asset_name'),'type_id':asset_type.id,'status_id':status.id,'model':required_text('model'),'serial':required_text('serial')}
            for key in ['brand','description']:
                if record.get(key) not in (None,''): data[key]=str(record[key]).strip()
            for key in ['received_date','handover_date']:
                if record.get(key):
                    value=record[key]
                    data[key]=value.date() if isinstance(value,datetime) else value if isinstance(value,date) else date.fromisoformat(str(value))
            asset=save(db,'assets',data,user); move(db,asset.id,Move(location_id=locations[0].id,reason='Imported asset registration'),user); created.append(asset)
        if not created: raise HTTPException(422,'The workbook has no asset rows')
        db.commit()
    except HTTPException:
        db.rollback(); raise
    except (ValueError,TypeError) as exc:
        db.rollback(); raise HTTPException(422,f'Invalid import value: {exc}')
    return {'created':len(created),'assets':[enriched(db,a) for a in created]}

@app.get('/api/tickets/{id}/detail')
def ticket_detail(id:int,user=Depends(current_user),db=Depends(get_db)):
    row=get(db,m.Ticket,id); data=enriched(db,row)
    stmt=select(m.TicketActivity).where(m.TicketActivity.ticket_id==id).order_by(m.TicketActivity.created_at)
    if user.role=='VIEWER': stmt=stmt.where(m.TicketActivity.internal==False)
    data['activities']=[enriched(db,r) for r in db.scalars(stmt)]
    data['maintenance']=[enriched(db,r) for r in db.scalars(select(m.Maintenance).where(m.Maintenance.ticket_id==id))]
    data['articles']=[enriched(db,get(db,m.Article,r.article_id)) for r in db.scalars(select(m.TicketArticle).where(m.TicketArticle.ticket_id==id))]
    data['attachments']=[serialize(r) for r in db.scalars(select(m.Attachment).where(m.Attachment.ticket_id==id))]
    return data

@app.post('/api/tickets/{id}/article',status_code=201)
def article_from_ticket(id:int,payload:dict,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'articles'); get(db,m.Ticket,id)
    data=parse_payload('articles',payload)
    article=save(db,'articles',data,user)
    save(db,'ticket-articles',{'ticket_id':id,'article_id':article.id},user)
    db.commit(); return enriched(db,article)

@app.post('/api/tickets/{id}/comments')
def comment(id:int,payload:Comment,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'tickets'); ticket=get(db,m.Ticket,id)
    activity(db,ticket,user,payload.body,'note' if payload.internal else 'comment',payload.internal)
    audit(db,user,'commented','tickets',ticket); db.commit(); return {'ok':True}

@app.post('/api/tickets/{id}/attachments')
async def attachment(id:int,file:UploadFile=File(...),user=Depends(current_user),db=Depends(get_db)):
    authorize(user,'tickets'); ticket=get(db,m.Ticket,id)
    content=await file.read(10*1024*1024+1)
    if len(content)>10*1024*1024: raise HTTPException(413,'Maximum file size is 10 MB')
    key=secrets.token_hex(24); directory=Path(settings.upload_dir); directory.mkdir(parents=True,exist_ok=True)
    path=directory/key; path.write_bytes(content)
    try:
        row=m.Attachment(ticket_id=id,user_id=user.id,filename=Path(file.filename or 'attachment').name,storage_key=key)
        db.add(row); db.flush(); activity(db,ticket,user,f'Attached {row.filename}','attachment'); audit(db,user,'uploaded','tickets',ticket); db.commit()
    except Exception: path.unlink(missing_ok=True); raise
    return serialize(row)

@app.get('/api/attachments/{id}')
def download(id:int,user=Depends(current_user),db=Depends(get_db)):
    row=get(db,m.Attachment,id); return FileResponse(Path(settings.upload_dir)/row.storage_key,filename=row.filename,media_type='application/octet-stream')

@app.get('/api/reports/{resource}/summary')
def report_summary(resource:str,q:str='',filters:str='{}',view:str='',location:str='',group_by:str='',user=Depends(current_user),db=Depends(get_db)):
    rows=[enriched(db,r) for r in db.scalars(query_rows(db,resource,user,q,filters,view,location))]
    totals={'Records':len(rows)}
    if resource=='tickets':
        durations=[(datetime.fromisoformat(r['resolved_at'])-datetime.fromisoformat(r['created_at'])).total_seconds()/3600 for r in rows if r['resolved_at']]
        totals['Avg resolution (hours)']=round(sum(durations)/len(durations),2) if durations else 0
        counts=Counter(r['asset_label'] for r in rows if r.get('asset_label'))
        totals['Assets with repeated issues']=sum(n>=3 for n in counts.values())
    if resource=='maintenance':
        totals['Total cost']=round(sum(r['cost'] or 0 for r in rows),2)
        totals['Active repairs']=sum(r['end_at'] is None for r in rows)
    if resource=='assignments': totals['Currently assigned']=sum(r['returned_at'] is None for r in rows)
    if resource=='switch-ports': totals['Connected ports']=sum(bool(r['connected_asset_id']) for r in rows)
    if resource=='ip-addresses': totals['Allocated IPs']=sum(bool(r['interface_id']) for r in rows)
    groups={}
    for row in rows:
        key=str(row.get(group_by) or 'Unspecified') if group_by else 'All records'
        group=groups.setdefault(key,{'group':key,'count':0,'cost':0})
        group['count']+=1; group['cost']+=float(row.get('cost') or 0)
    return {'totals':totals,'groups':list(groups.values())}

@app.get('/api/reports/{resource}/export')
def export(resource:str,format:str='csv',q:str='',filters:str='{}',view:str='',location:str='',user=Depends(current_user),db=Depends(get_db)):
    stmt=query_rows(db,resource,user,q,filters,view,location)
    rows=[enriched(db,r) for r in db.scalars(stmt)]
    columns=list(rows[0]) if rows else [c.name for c in inspect(RESOURCES[resource]).columns if c.name!='password_hash']
    def safe(v):
        value=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else '' if v is None else str(v)
        return "'"+value if value.startswith(('=','+','-','@','\t','\r')) else value
    if format=='xlsx':
        book=Workbook(); sheet=book.active; sheet.append(columns)
        for row in rows: sheet.append([safe(row.get(c)) for c in columns])
        buffer=io.BytesIO(); book.save(buffer); buffer.seek(0); mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif format=='csv':
        out=io.StringIO(); writer=csv.writer(out); writer.writerow(columns)
        for row in rows: writer.writerow([safe(row.get(c)) for c in columns])
        buffer=io.BytesIO(('\ufeff'+out.getvalue()).encode()); mime='text/csv; charset=utf-8'
    else: raise HTTPException(422,'Use csv or xlsx')
    return StreamingResponse(buffer,media_type=mime,headers={'Content-Disposition':f'attachment; filename="{resource}.{format}"'})

@app.get('/api/{resource}')
def list_resource(resource:str,q:str='',filters:str='{}',page:int=Query(1,ge=1),page_size:int=Query(20,ge=1,le=100),sort:str='id',direction:str='desc',view:str='',location:str='',user=Depends(current_user),db=Depends(get_db)):
    stmt=query_rows(db,resource,user,q,filters,view,location); model=RESOURCES[resource]
    if sort=='password_hash': raise HTTPException(422,'Invalid sort')
    if sort not in inspect(model).columns:
        rows=[enriched(db,r) for r in db.scalars(stmt)]
        display=sort.removesuffix('_id')+'_label' if sort.endswith('_id') else sort
        if rows and display not in rows[0]: raise HTTPException(422,'Invalid sort')
        rows.sort(key=lambda r:str(r.get(display) or ''),reverse=direction=='desc')
        return {'items':rows[(page-1)*page_size:page*page_size],'total':len(rows),'page':page,'page_size':page_size}
    total=db.scalar(select(func.count()).select_from(stmt.subquery()))
    col=getattr(model,sort); stmt=stmt.order_by(col.desc() if direction=='desc' else col.asc())
    return {'items':[enriched(db,r) for r in db.scalars(stmt.offset((page-1)*page_size).limit(page_size))],'total':total,'page':page,'page_size':page_size}

@app.get('/api/{resource}/{id}')
def detail(resource:str,id:int,user=Depends(current_user),db=Depends(get_db)):
    if resource not in RESOURCES: raise HTTPException(404)
    if resource=='users' and user.role!='ADMIN': raise HTTPException(403)
    row=get(db,RESOURCES[resource],id)
    if resource=='ticket-activities' and row.internal and user.role=='VIEWER': raise HTTPException(403)
    return enriched(db,row)

def parse_payload(resource,payload,partial=False):
    if resource not in SCHEMAS: raise HTTPException(405,'Resource is read-only')
    try: return SCHEMAS[resource][int(partial)].model_validate(payload).model_dump(exclude_unset=True)
    except ValidationError as e: raise HTTPException(422,jsonable_encoder(e.errors(),custom_encoder={ValueError:str}))

@app.post('/api/{resource}',status_code=201)
def create(resource:str,payload:dict,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,resource)
    if resource=='assets': raise HTTPException(422,'Use asset registration and provide an initial location')
    data=parse_payload(resource,payload); obj=save(db,resource,data,user); db.commit(); return enriched(db,obj)

@app.patch('/api/{resource}/{id}')
def edit(resource:str,id:int,payload:dict,user=Depends(current_user),db=Depends(get_db)):
    authorize(user,resource); data=parse_payload(resource,payload,True); obj=get(db,RESOURCES[resource],id)
    if resource=='users' and obj.id==user.id and data.get('role',user.role)!=user.role: raise HTTPException(422,'Cannot change your own role')
    obj=save(db,resource,data,user,obj); db.commit(); return enriched(db,obj)

@app.delete('/api/{resource}/{id}')
def archive(resource:str,id:int,user=Depends(current_user),db=Depends(get_db)):
    if user.role!='ADMIN': raise HTTPException(403,'Admin only')
    if resource not in SCHEMAS or resource in {'tickets','maintenance','ticket-articles'}: raise HTTPException(405,'Historical records cannot be archived')
    obj=get(db,RESOURCES[resource],id)
    if resource=='vlans' and any(id in p.tagged_vlans for p in db.scalars(select(m.SwitchPort))): raise HTTPException(409,'VLAN is referenced by tagged switch ports')
    if resource=='master-data':
        from .defaults import DEFAULT_MASTER_DATA
        if obj.code in [name.lower().replace(' ','_') for name in DEFAULT_MASTER_DATA.get(obj.group,[])]: raise HTTPException(422,'Built-in workflow identifiers cannot be archived; change their display names instead')
    if resource=='users' and obj.id==user.id: raise HTTPException(422,'Cannot archive yourself')
    # Prevent archiving records still referenced by other records.
    for model in RESOURCES.values():
        for c in inspect(model).columns:
            if any(fk.column.table.name==obj.__tablename__ for fk in c.foreign_keys):
                if db.scalar(select(func.count()).select_from(model).where(c==id)): raise HTTPException(409,'Record is referenced; retain it for history')
    old=serialize(obj); obj.archived=True; audit(db,user,'archived',resource,obj,old); db.commit(); return {'ok':True}
